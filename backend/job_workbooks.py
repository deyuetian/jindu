"""Read recruitment workbooks without Excel, macros, external links or formula execution."""
import base64
import io
import json
import re
import zipfile
from datetime import date, datetime
from xml.etree import ElementTree

MAX_ROWS = 12000
MAX_CELLS = 300000
MAX_CHARS = 1300000


def cell_text(value, number_format=''):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if int(value) == value:
            number = str(int(value))
            if re.fullmatch('0{2,20}', number_format):
                return number.zfill(len(number_format))
            return number
    return str(value).replace('\r\n', '\n').strip()


CATALOG = re.compile(r'专业.{0,8}目录|学科.{0,8}目录|(?:本科|研究生|硕士|博士|职业教育).{0,8}目录|自主设置.{0,8}学科名单')

def table_from_rows(name, rows, merges, warnings, source_name=''):
    # Only actual merged ranges are propagated. Blank ordinary cells stay blank.
    values = {i: list(cells) for i, cells in rows}
    for first_row, last_row, first_col, last_col in merges:
        if (last_row-first_row+1)*(last_col-first_col+1)>50000:
            warnings.append(name+'：合并区域过大，未展开'); continue
        value = values.get(first_row, [])
        value = value[first_col-1] if len(value)>=first_col else ''
        for row in range(first_row, last_row+1):
            if row not in values: continue
            for col in range(first_col, min(last_col, 100)+1):
                while len(values[row])<col: values[row].append('')
                values[row][col-1] = value
    ordered = sorted(values.items())
    heading=source_name+' '+name+' '+' '.join(' '.join(dict.fromkeys(cells)) for _,cells in ordered[:4])
    # A university/discipline catalogue is reference data even when it has a unit column.
    # Explicit recruitment headers win for real job sheets inside mixed archives.
    job_header=any(any(re.fullmatch(r'(?:招聘|招录|拟聘|报考)?(?:岗位|职位)(?:名称|代码|编号|类别|序号)?|招聘人数|招录人数|招聘计划|招考人数|拟聘人数',re.sub(r'\s','',v)) for v in cells) for _,cells in ordered[:40])
    reference=bool(CATALOG.search(heading)) and not job_header
    if reference:
        return {'sheet':name,'sourceName':source_name,'headers':[],'rows':[],'notes':ordered,'complete':True,'kind':'reference'}
    header = None
    for i, cells in ordered[:40]:
        labels = ''.join(cells)
        if re.search('专业|学科', labels) and re.search('岗位|职位|职务|单位|部门', labels):
            # A merged title alone does not establish a table header.
            if len(set(v for v in cells if v))>=2:
                header = i; break
    if header is None:
        return {'sheet':name,'sourceName':source_name,'headers':[],'rows':[],'notes':ordered,'complete':False,'kind':'unknown'}
    headers = values[header][:]
    last_header = header
    # Degree subcolumns commonly occupy the next one or two header rows.
    for index in range(header+1, header+3):
        cells = values.get(index, [])
        filled = [v for v in cells if v]
        if not filled or not any(re.fullmatch('本科|研究生|硕士|博士|专科|专业要求|专业名称|专业代码', v) for v in filled): break
        if any(re.search('[0-9]{4,}|土壤学|工程师|科员|管理岗|技术岗', v) for v in filled): break
        if any(len(v)>30 for v in filled): break
        for col, v in enumerate(cells):
            while len(headers)<=col: headers.append('')
            if v and v not in headers[col]: headers[col] = (headers[col]+' / '+v).strip(' /')
        last_header = index
    headers = [v or f'未命名列 {i+1}' for i,v in enumerate(headers)]
    data=[]; notes=[]
    for index, cells in ordered:
        if index<=last_header:
            if index<header: notes.append([index,cells])
            continue
        if not any(cells): continue
        nonempty=[v for v in cells if v]
        if len(set(nonempty))==1 and (nonempty[0].startswith(('注','说明','备注：')) or len(nonempty)==1):
            notes.append([index,cells]);continue
        if cells==values[header]: continue
        fields=[{'column':i+1,'header':headers[i] if i<len(headers) else f'未命名列 {i+1}','value':v} for i,v in enumerate(cells) if v]
        text=f'[{name} 第 {index} 行] '+' | '.join(f'【{f["header"]}】{f["value"]}' for f in fields)
        data.append({'row':index,'fields':fields,'text':text})
    return {'sheet':name,'sourceName':source_name,'headers':headers,'rows':data,'notes':notes,'complete':True,'kind':'positions'}


def read_workbook(path, extension, source_name=''):
    warnings=[]; tables=[]; count=0; chars=0
    if extension=='.xls':
        import xlrd
        book=xlrd.open_workbook(file_contents=path.read_bytes(), formatting_info=True, on_demand=True)
        try:
            for sheet in book.sheets():
                rows=[]
                for i in range(min(sheet.nrows,MAX_ROWS)):
                    cells=[]
                    for col in range(min(sheet.ncols,100)):
                        cell=sheet.cell(i,col); fmt=''
                        try: fmt=book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                        except (IndexError,KeyError): pass
                        value=xlrd.xldate_as_datetime(cell.value,book.datemode) if cell.ctype==3 else cell.value
                        cells.append(cell_text(value,fmt))
                    count+=len(cells);chars+=sum(map(len,cells))
                    if count>MAX_CELLS or chars>MAX_CHARS: warnings.append('工作簿超过读取上限，剩余行待核对');break
                    rows.append((i+1,cells))
                if sheet.nrows>MAX_ROWS or sheet.ncols>100:warnings.append(sheet.name+'：超出行列读取上限')
                merges=[(r+1,re,c+1,ce) for r,re,c,ce in sheet.merged_cells]
                tables.append(table_from_rows(sheet.name,rows,merges,warnings,source_name))
                if count>MAX_CELLS or chars>MAX_CHARS:break
        finally:book.release_resources()
    else:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
        with zipfile.ZipFile(path) as archive:
            if len(archive.infolist())>20000 or sum(v.file_size for v in archive.infolist())>150000000:
                raise ValueError('Workbook expansion limit')
            with path.open('rb') as handle:
                book=load_workbook(handle,read_only=True,data_only=True,keep_links=False)
                try:
                    for sheet in book:
                        merges=[]; formula=False
                        # Streaming the XML avoids loading a second in-memory workbook.
                        with archive.open(sheet._worksheet_path.lstrip('/')) as xml:
                            for _,element in ElementTree.iterparse(xml,events=('end',)):
                                if element.tag.endswith('}mergeCell'):
                                    c,r,ce,re_=range_boundaries(element.attrib['ref']);merges.append((r,re_,c,ce))
                                elif element.tag.endswith('}f'):formula=True
                                element.clear()
                        if formula:warnings.append(sheet.title+'：含公式，仅读取保存的计算结果，需核对')
                        rows=[]
                        for index,row in enumerate(sheet.iter_rows(max_col=min(sheet.max_column or 100,100)),1):
                            if index>MAX_ROWS:warnings.append(sheet.title+'：超出行读取上限');break
                            cells=[cell_text(c.value,c.number_format or '') for c in row]
                            count+=len(cells);chars+=sum(map(len,cells))
                            if count>MAX_CELLS or chars>MAX_CHARS:warnings.append('工作簿超过读取上限，剩余行待核对');break
                            rows.append((index,cells))
                        if (sheet.max_column or 0)>100:warnings.append(sheet.title+'：超出列读取上限')
                        tables.append(table_from_rows(sheet.title,rows,merges,warnings,source_name))
                        if count>MAX_CELLS or chars>MAX_CHARS:break
                finally:book.close()
    text=[]
    for table in tables:
        if not table['complete']:warnings.append(table['sheet']+'：未识别岗位表头，保留正文待分析')
        for number,cells in table['notes']:text.append(f'[{table["sheet"]} 第 {number} 行] '+' | '.join(cells))
        if table['headers']:text.append(f'[{table["sheet"]} 表头] '+' | '.join(table['headers']))
        text.extend(row['text'] for row in table['rows'])
    result={'text':'\n'.join(text),'links':[],'error':'；'.join(dict.fromkeys(warnings)),'tables':tables,'complete':not warnings,'parserVersion':3}
    # JSON output is bounded independently from source bytes.
    if len(json.dumps(result,ensure_ascii=False))>4800000:
        return {'text':result['text'][:MAX_CHARS],'links':[],'tables':[],'complete':False,'error':'工作簿结构超过读取上限，未完成筛选','parserVersion':3}
    return result


def export_jobs(payload):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    book=Workbook();book.remove(book.active)
    labels=['已核查条件满足','待补充或核对','有不满足条件']
    headers=['岗位代码','岗位','单位','结论','匹配学历','专业依据','其他资格','材料与步骤','来源附件','工作表','原始行号','原文链接','工作地点','原表要求']
    jobs=payload.get('jobs',[])
    if len(jobs)>3000:raise ValueError('Export limit')
    def literal(value):
        text=str(value if value is not None else '')[:32000]
        # Excel must never interpret scraped text as a formula.
        return "'"+text if text.lstrip().startswith(('=','+','-','@','\t','\r')) else text
    for label in labels:
        sheet=book.create_sheet(label);sheet.append(headers)
        for job in jobs:
            verdict=labels[0] if job.get('match')=='符合' else labels[2] if job.get('match')=='不符合' else labels[1]
            if verdict!=label:continue
            ref=job.get('sourceRow') or {}
            checks='\n'.join(f'{g["label"]}：'+ '; '.join(f'{i["label"]} / {i["status"]} / {i["actual"]}' for i in g['items']) for g in (job.get('eligibilityChecks') or {}).get('groups',[]))
            materials='\n'.join(f'{i["label"]}：{i.get("detail","")}（{i.get("phase","未说明")} / {i.get("method","")}）' for i in job.get('applicationItems',[]))
            row=[job.get('positionCode'),job.get('title'),job.get('organization'),verdict,job.get('matchedLevel'),job.get('majorReason'),checks,materials,ref.get('fileName'),ref.get('sheet'),ref.get('row'),job.get('url'),job.get('displayRegion') or job.get('province'), '\n'.join(str(f.get('header',''))+'：'+str(f.get('value','')) for f in (job.get('sourcePosition') or {}).get('fields',[]))]
            sheet.append([literal(v) for v in row])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245D50')
        for index in range(1,len(headers)+1):sheet.column_dimensions[get_column_letter(index)].width=28 if index in [6,7,8] else 18
        for row in sheet.iter_rows(min_row=2):
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    out=io.BytesIO();book.save(out)
    return {'base64':base64.b64encode(out.getvalue()).decode('ascii')}
