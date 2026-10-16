"""Real workbook fixtures and parser/export checks; no network or AI requests."""
from pathlib import Path
import base64, importlib.util, json, sys, zipfile
from openpyxl import Workbook, load_workbook

root=Path(__file__).resolve().parents[1];out=root/'.check/v10';out.mkdir(parents=True,exist_ok=True)
spec=importlib.util.spec_from_file_location('job_workbooks',root/'backend/job_workbooks.py');module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
book=Workbook();sheet=book.active;sheet.title='省直选调'
sheet.append(['2027年度选调岗位表']);sheet.merge_cells('A1:H1')
sheet.append(['单位名称','岗位名称','岗位代码','报考专业',None,'政治面貌','经历要求','提交材料'])
sheet.append([None,None,None,'硕士','本科',None,None,None]);sheet.merge_cells('D2:E2')
for col in ['A','B','C','F','G','H']:sheet.merge_cells(f'{col}2:{col}3')
sheet.append(['自然资源研究所','专业技术',1,'农学','土地资源管理','中共党员','无需工作经历','报名在线提交推荐表（PDF）'])
sheet.append([None,'专业技术',2,'农业资源与环境','土地资源管理','群众','土壤调查工作满12个月，截至2026年9月1日','资格复审提交毕业证复印件'])
sheet.merge_cells('A4:A5')
sheet.append(['信息研究所','专业技术',3,'计算机科学与技术','计算机科学与技术','不限','无',''])
sheet['C4'].number_format='000';sheet['C5'].number_format='000';sheet['C6'].number_format='000'
other=book.create_sheet('市直选调');other.append(['岗位代码','职位名称','单位名称','学历','专业要求','备注']);other.append(['004','土壤研究','农业研究所','硕士','土壤学','党员优先'])
book.save(out/'positions.xlsx')
parsed=module.read_workbook(out/'positions.xlsx','.xlsx')
assert parsed['complete'],parsed['error']
rows=parsed['tables'][0]['rows'];assert len(rows)==3
assert rows[0]['row']==4 and rows[0]['fields'][2]['value']=='001'
assert rows[1]['fields'][0]['value']=='自然资源研究所'
assert rows[0]['fields'][3]['header']=='报考专业 / 硕士'
assert parsed['tables'][1]['rows'][0]['row']==2
assert '【报考专业 / 硕士】农学' in parsed['text']
formula=Workbook();s=formula.active;s.append(['岗位名称','专业要求']);s.append(['岗位','=HYPERLINK("https://example.invalid","土壤学")']);formula.save(out/'formula.xlsx')
unsafe=module.read_workbook(out/'formula.xlsx','.xlsx');assert not unsafe['complete'] and '公式' in unsafe['error']
with zipfile.ZipFile(out/'positions.zip','w') as z:
    z.write(out/'positions.xlsx','附件1/岗位.xlsx');z.writestr('../说明.txt','报名须提交推荐表。')
sys.path.insert(0,str(root/'backend'))
spec=importlib.util.spec_from_file_location('file_tools',root/'backend/file-tools.py');files=importlib.util.module_from_spec(spec);spec.loader.exec_module(files)
z=files.parse(out/'positions.zip','positions.zip');assert '附件1/岗位.xlsx / 省直选调' in z['text'];assert '推荐表' in z['extraText'];assert len(z['tables'])==2
jobs=[{'title':'=HYPERLINK("https://example.invalid")','match':'符合','positionCode':'001','sourceRow':{'sheet':'省直选调','row':4,'fileName':'positions.xlsx'},'applicationItems':[]},{'title':'待查岗','match':'待核对'},{'title':'不符岗','match':'不符合'}]
data=module.export_jobs({'jobs':jobs});(out/'筛选结果.xlsx').write_bytes(base64.b64decode(data['base64']))
export=load_workbook(out/'筛选结果.xlsx');assert len(export.worksheets)==3
assert export.worksheets[0]['B2'].data_type=='s' and export.worksheets[0]['B2'].value.startswith("'")
assert all(s.auto_filter.ref and s.max_row==2 for s in export)
(out/'workbook-fixture.json').write_text(json.dumps(parsed,ensure_ascii=False),encoding='utf-8')
(out/'workbook-check.json').write_text(json.dumps({'passed':True,'mergedHeaders':True,'mergedValues':True,'leadingZeros':True,'multiSheet':True,'formulaNotExecuted':True,'zipSafePaths':True,'actualFilteredExport':True}),encoding='utf-8')
print('Workbook extraction, ZIP attachments and filtered Excel export passed.')
