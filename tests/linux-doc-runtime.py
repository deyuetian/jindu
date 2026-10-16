"""Convert a generated DOCX to DOC, then read it using the image's parser.

Run in the Linux image with /tmp writable. No downloaded or personal documents.
"""
from pathlib import Path
import hashlib
import subprocess
import sys
import tempfile
from docx import Document

sys.path.insert(0, '/app/backend')
from legacy_doc import legacy_doc_text

with tempfile.TemporaryDirectory(prefix='jindu-doc-runtime-') as directory:
    base = Path(directory)
    source = base / 'synthetic.docx'
    document = Document()
    document.add_paragraph('JINDU_SYNTHETIC_DOCUMENT')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Position'
    table.cell(0, 1).text = 'Requirement'
    table.cell(1, 0).text = 'Synthetic analyst'
    table.cell(1, 1).text = 'Synthetic qualification'
    document.save(source)
    subprocess.run(['soffice', '-env:UserInstallation=' + (base/'profile').as_uri(),
                    '--headless', '--convert-to', 'doc:MS Word 97', '--outdir',
                    str(base), str(source)], check=True, capture_output=True,
                   timeout=60)
    legacy = base / 'synthetic.doc'
    before = hashlib.sha256(legacy.read_bytes()).hexdigest()
    text = legacy_doc_text(legacy)
    assert 'JINDU_SYNTHETIC_DOCUMENT' in text
    assert 'Synthetic analyst' in text and 'Synthetic qualification' in text
    assert hashlib.sha256(legacy.read_bytes()).hexdigest() == before
    print('{"passed":true,"legacyDOC":true,"table":true,"originalPreserved":true,"synthetic":true}')
