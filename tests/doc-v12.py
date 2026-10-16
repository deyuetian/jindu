"""Legacy document conversion tests. Uses an existing office runtime, never installs one."""
from pathlib import Path
import hashlib,importlib.util,json,os,sys,tempfile,unittest,shutil
from unittest.mock import patch
from docx import Document
root=Path(__file__).resolve().parents[1];out=root/'.check/v12/doc-tests';out.mkdir(parents=True,exist_ok=True)
tempfile.tempdir=str(out);os.environ['TEMP']=str(out);os.environ['TMP']=str(out)
sys.path.insert(0,str(root/'backend'))
import legacy_doc
spec=importlib.util.spec_from_file_location('file_tools',root/'backend/file-tools.py');file_tools=importlib.util.module_from_spec(spec);spec.loader.exec_module(file_tools)

class DocumentReading(unittest.TestCase):
    def setUp(self):
        self.temporary=tempfile.TemporaryDirectory(prefix='doc-test-',dir=out);self.base=Path(self.temporary.name)
    def tearDown(self):self.temporary.cleanup()
    def make_docx(self):
        document=Document();document.add_paragraph('前文报名条件');table=document.add_table(rows=2,cols=2)
        table.cell(0,0).text='学校';table.cell(0,1).text='专业';table.cell(1,0).text='示例大学';table.cell(1,1).text='土壤学';document.add_paragraph('后文提交材料')
        path=self.base/'source.docx';document.save(path);return path
    def test_preserves_paragraph_table_order_and_row_numbers(self):
        path=self.make_docx();text=legacy_doc.read_docx(path)
        self.assertLess(text.index('前文报名条件'),text.index('[表 1 第 1 行]'));self.assertLess(text.index('示例大学 | 土壤学'),text.index('后文提交材料'))
    def test_reads_docx_with_doc_suffix_without_converter(self):
        path=self.make_docx();copy=self.base/'disguised.doc';copy.write_bytes(path.read_bytes())
        with patch.object(legacy_doc.shutil,'which',return_value=None):self.assertIn('土壤学',legacy_doc.legacy_doc_text(copy))
    def test_retains_empty_documents_as_unreadable(self):
        path=self.base/'empty.docx';Document().save(path);self.assertIn('未读取到文字',file_tools.parse(path,'empty.docx')['error'])
    def test_missing_converter_returns_actionable_error(self):
        path=self.base/'old.doc';path.write_bytes(b'legacy')
        with patch.dict(os.environ,{'JINDU_SOFFICE':''}),patch.object(legacy_doc.shutil,'which',return_value=None):
            with self.assertRaisesRegex(legacy_doc.DocReadError,'DOC 转换组件'):legacy_doc.legacy_doc_text(path)
        self.assertEqual(path.read_bytes(),b'legacy')
    def test_failed_conversion_preserves_original_and_private_profile(self):
        path=self.base/'broken.doc';path.write_bytes(b'legacy');commands=[]
        class Process:
            returncode=1
            def __init__(self,command,**kwargs):
                from urllib.parse import unquote,urlparse
                import xml.etree.ElementTree as ET
                commands.append(command);uri=next(c.split('=',1)[1] for c in command if c.startswith('-env:UserInstallation='))
                location=unquote(urlparse(uri).path).lstrip('/') if os.name=='nt' else unquote(urlparse(uri).path)
                xml=ET.parse(Path(location)/'user/registrymodifications.xcu');values={p.attrib['{http://openoffice.org/2001/registry}name']:p.findtext('value') for p in xml.findall('.//prop')}
                assert values['Link']=='2' and values['DisableMacrosExecution']=='true' and values['DisableActiveContent']=='true'
            def wait(self,**kwargs):return 1
            def poll(self):return 1
        with patch.dict(os.environ,{'JINDU_SOFFICE':'fixture-converter'}),patch.object(legacy_doc.subprocess,'Popen',Process):
            with self.assertRaisesRegex(legacy_doc.DocReadError,'损坏或受密码保护'):legacy_doc.legacy_doc_text(path)
        self.assertEqual(path.read_bytes(),b'legacy');self.assertIn('--headless',commands[0]);self.assertTrue(any(a.startswith('-env:UserInstallation=') for a in commands[0]))
    def test_zip_docx_retains_table_evidence(self):
        import zipfile
        path=self.make_docx();archive=self.base/'attachments.zip'
        with zipfile.ZipFile(archive,'w') as z:z.write(path,'材料.docx')
        result=file_tools.parse(archive,'材料.zip');self.assertFalse(result['error']);self.assertIn('[表 1 第 2 行] 示例大学 | 土壤学',result['text'])

if __name__=='__main__':
    suite=unittest.defaultTestLoader.loadTestsFromTestCase(DocumentReading);result=unittest.TextTestRunner(verbosity=1).run(suite)
    report={'unitTests':result.testsRun,'unitPassed':result.wasSuccessful(),'realDocuments':[]}
    fixtures=Path(os.environ.get('JINDU_DOC_FIXTURES',str(root/'.check/v12/doc-fixtures')))
    if (fixtures/'manifest.json').exists():
        if os.name=='nt':
            wrapper=Path(os.environ.get('JINDU_SOFFICE') or shutil.which('soffice.exe') or 'missing-wrapper');assert wrapper.is_file() and 'LibreOffice/program' not in wrapper.as_posix(),'Existing compatible wrapper unavailable; no installation attempted'
            os.environ['JINDU_SOFFICE']=str(wrapper)
        for meta in json.loads((fixtures/'manifest.json').read_text(encoding='utf-8')):
            path=fixtures/(meta['id']+'.doc');before=hashlib.sha256(path.read_bytes()).hexdigest();assert before==meta['sha256']
            parsed=file_tools.parse(path,meta['name']);assert not parsed['error'],parsed['error'];text=parsed['text'];assert len(text)>100
            expected=['中国农业大学'] if '范围' in meta['name'] else ['姓名','政治面貌']
            assert all(term in text for term in expected),meta['name']
            assert hashlib.sha256(path.read_bytes()).hexdigest()==before
            (out/(meta['id']+'.txt')).write_text(text,encoding='utf-8')
            report['realDocuments'].append({'name':meta['name'],'characters':len(text),'tableRows':text.count('[表 '),'expectedFieldsPresent':True,'originalUnchanged':True})
    report['passed']=report['unitPassed'];(out/'result.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False));raise SystemExit(0 if report['passed'] else 1)
